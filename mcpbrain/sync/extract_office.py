"""DOCX and PPTX -> Blocks.

Both walk their document tree in reading order rather than concatenating text
runs, so headings, tables and speaker notes survive as structured Blocks
(2026-09-24 extraction-fidelity spec) instead of being flattened.
"""
import io
import logging

from mcpbrain.sync.blocks import Heading, Paragraph, PartialBlocks, TableBlock

log = logging.getLogger(__name__)


def _docx_heading_level(style_name: str) -> int:
    name = (style_name or "").strip()
    if name == "Title":
        return 1
    if name.startswith("Heading"):
        tail = name.split()[-1]
        return int(tail) if tail.isdigit() else 1
    return 0


def _docx_table_rows(table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        seen, cells = set(), []
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                continue                     # a merged cell spans several grid slots
            seen.add(key)
            cells.append(cell.text.strip())
        rows.append(cells)
    return rows


def extract_blocks_from_docx(content_bytes: bytes) -> list:
    """DOCX -> Blocks walking the body IN ORDER (w:p and w:tbl), so tables sit
    where they appear. Heading/Title styles become Headings; text boxes follow
    their anchor paragraph; section headers first and footers last, each once."""
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph as DocxParagraph
        doc = Document(io.BytesIO(content_bytes))
    except Exception as exc:
        log.warning("docx: open failed: %s", exc)
        return []
    out: list = []
    try:
        headers, footers = [], []
        for section in doc.sections:
            for part, bucket in ((section.header, headers), (section.footer, footers)):
                text = "\n".join(p.text for p in part.paragraphs if p.text.strip()).strip()
                if text and text not in bucket:
                    bucket.append(text)
        out.extend(Paragraph(h) for h in headers)
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                p = DocxParagraph(child, doc)
                text = p.text.strip()
                if text:
                    level = _docx_heading_level(p.style.name if p.style is not None else "")
                    out.append(Heading(level, text) if level else Paragraph(text))
                for box in child.iter(qn("w:txbxContent")):
                    box_text = "\n".join(
                        "".join(t.text or "" for t in bp.iter(qn("w:t")))
                        for bp in box.iter(qn("w:p"))).strip()
                    if box_text:
                        out.append(Paragraph(box_text))
            elif child.tag == qn("w:tbl"):
                rows = _docx_table_rows(DocxTable(child, doc))
                if any(any(c for c in r) for r in rows):
                    out.append(TableBlock(rows))
        out.extend(Paragraph(f) for f in footers)
        return out
    except Exception as exc:
        log.warning("docx: extraction failed after %d blocks: %s", len(out), exc)
        return PartialBlocks(out) if out else []


def _pptx_walk(shapes, skip_id, out: list) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _pptx_walk(shape.shapes, skip_id, out)
            continue
        if shape.shape_id == skip_id:
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            lines = ["".join(r.text for r in p.runs).strip()
                     for p in shape.text_frame.paragraphs]
            text = "\n".join(line for line in lines if line)
            if text:
                out.append(Paragraph(text))
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [[c.text.strip() for c in row.cells] for row in shape.table.rows]
            if any(any(c for c in r) for r in rows):
                out.append(TableBlock(rows))


def extract_blocks_from_pptx(content_bytes: bytes) -> list:
    """PPTX -> Blocks: a Heading per slide (rendered 'Slide N: <title>', the
    'Slide N: ' part a label), shapes walked recursively through groups, tables
    as TableBlocks, and speaker notes as a trailing Paragraph labelled
    'Notes: '. Labels are rendered but are never source spans."""
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(content_bytes))
    except Exception as exc:
        log.warning("pptx: presentation open failed: %s", exc)
        return []
    out: list = []
    try:
        for n, slide in enumerate(prs.slides, start=1):
            title_shape = slide.shapes.title
            title = title_shape.text_frame.text.strip() if title_shape is not None else ""
            # "Slide N: " is synthesised, not slide text: a label, never a span
            out.append(Heading(2, title, label=f"Slide {n}: ") if title
                       else Heading(2, "", label=f"Slide {n}"))
            _pptx_walk(slide.shapes,
                       title_shape.shape_id if title_shape is not None else None, out)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    out.append(Paragraph(notes, label="Notes: "))
        return out
    except Exception as exc:
        log.warning("pptx: extraction failed after %d blocks: %s", len(out), exc)
        return PartialBlocks(out) if out else []
