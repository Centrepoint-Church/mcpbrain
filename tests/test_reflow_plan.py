from mcpbrain.reflow import lineage_key, norm, plan, stitch
from mcpbrain.sync.normalise import Chunk


def _old(i, text, total, enriched=1, state=None, fid="F", version=3):
    return {"doc_id": f"gdrive-{fid}-{i}", "text": text,
            "metadata": {"source_type": "gdrive", "file_id": fid, "chunk_index": i,
                         "chunk_total": total},
            "enriched": enriched, "enriched_version": version, "enrich_state": state,
            "salience": 0.5, "memory_tier": "warm", "memory_type": "semantic"}


def _new(i, text, spans=None, fid="F"):
    return Chunk(f"gdrive-{fid}-{i}", text, "h", {"source_type": "gdrive", "file_id": fid,
                 "chunk_index": i}, spans or [text])


def test_stitch_removes_word_overlap():
    a = "one two three four five six"
    b = "five six seven eight"
    o, offs = stitch([a, b])
    assert o == "one two three four five six seven eight"
    # Offsets partition the stitched text: each input owns what it CONTRIBUTED,
    # i.e. its normalised text minus the overlap prefix the seam dropped.
    assert offs[0] == (0, len(a))
    assert offs[1] == (o.index("seven"), len(o))
    assert o[offs[0][0]:offs[0][1]] == norm(a)
    assert o[offs[1][0]:offs[1][1]] == norm(b)[len("five six "):]


def test_stitch_offsets_invariant_on_messy_whitespace():
    texts = ["  alpha\n\nbeta  gamma\tdelta ", "gamma delta epsilon", "zeta", "", "eta theta"]
    o, offs = stitch(texts)
    assert o == "alpha beta gamma delta epsilon zeta eta theta"
    dropped = [0, len("gamma delta "), 0, 0, 0]
    for t, (s, e), d in zip(texts, offs, dropped):
        assert o[s:e] == norm(t)[d:]
    # a partition: spans are ordered and never overlap
    assert all(offs[i][1] <= offs[i + 1][0] for i in range(len(offs) - 1))


def test_stitch_whole_chunk_inside_overlap_gets_empty_span():
    o, offs = stitch(["a b c d", "c d"])
    assert o == "a b c d"
    assert offs[1] == (len(o), len(o))


def test_lineage_keys():
    assert lineage_key("gdrive-F-3", {"source_type": "gdrive"}) == "gdrive-F"
    assert lineage_key("gmail-M-att-2-0", {"source_type": "gmail"}) == "gmail-M-att-2"
    assert lineage_key("cal-E", {"source_type": "calendar", "event_id": "E"}) == "cal-E"
    assert lineage_key("cal-E-1", {"source_type": "calendar", "event_id": "E"}) == "cal-E"


def test_full_coverage_inherits_and_remaps():
    old = [_old(0, "alpha beta gamma delta", 2), _old(1, "gamma delta epsilon zeta", 2)]
    new = [_new(0, "alpha beta\ngamma", ["alpha beta\ngamma"]),
           _new(1, "delta epsilon zeta", ["delta epsilon zeta"])]
    p = plan(old, new)
    assert p.content_equal
    assert all(r.covered and r.enriched == 1 for r in p.rows)
    assert p.remap["gdrive-F-0"] == "gdrive-F-0"
    assert p.remap["gdrive-F-1"] in {"gdrive-F-0", "gdrive-F-1"}
    assert p.deletes == []


def test_identical_overlapping_chunks_remap_to_identity():
    """chunk_text word-splits with a 50-word overlap. An unchanged re-chunk with
    the same boundaries must map i -> i, not onto the chunk whose tail holds the
    overlap (an inclusive start offset would map every i onto i-1)."""
    w = [f"w{n}" for n in range(30)]
    texts = [" ".join(w[0:12]), " ".join(w[8:22]), " ".join(w[18:30])]
    old = [_old(i, t, 3) for i, t in enumerate(texts)]
    p = plan(old, [_new(i, t) for i, t in enumerate(texts)])
    assert p.remap == {f"gdrive-F-{i}": f"gdrive-F-{i}" for i in range(3)}
    assert set(p.reasons.values()) == {"exact"}
    assert p.content_equal and all(r.covered and r.enriched for r in p.rows)


def test_new_text_is_uncovered_and_not_inherited():
    old = [_old(0, "alpha beta", 1)]
    new = [_new(0, "alpha beta"), _new(1, "Notes: brand new speaker notes")]
    p = plan(old, new)
    assert p.rows[0].covered and not p.rows[1].covered
    assert p.rows[1].enriched == 0 and p.rows[1].enrich_state is None
    assert p.rows[1].salience is None and p.rows[1].memory_tier is None
    assert not p.content_equal
    assert p.unequal == ["gdrive-F"]


def test_partial_word_is_not_coverage():
    """'pha bet' occurs inside 'alpha beta' as characters but not as words."""
    p = plan([_old(0, "alpha beta", 1)], [_new(0, "pha bet")])
    assert not p.rows[0].covered


def test_unenriched_old_means_not_enriched_new():
    old = [_old(0, "alpha beta", 1, enriched=0)]
    p = plan(old, [_new(0, "alpha beta")])
    assert p.rows[0].covered and p.rows[0].enriched == 0
    assert p.rows[0].enriched_version == 0


def test_mixed_enrichment_is_not_enriched():
    old = [_old(0, "alpha beta", 2), _old(1, "gamma delta", 2, enriched=0)]
    p = plan(old, [_new(0, "alpha beta gamma delta")])
    assert p.rows[0].covered and p.rows[0].enriched == 0


def test_enriched_version_is_the_oldest_overlapped():
    old = [_old(0, "alpha beta", 2, version=3), _old(1, "gamma delta", 2, version=2)]
    p = plan(old, [_new(0, "alpha beta gamma delta")])
    assert p.rows[0].enriched == 1 and p.rows[0].enriched_version == 2


def test_fewer_new_chunks_deletes_and_maps_old_ids():
    old = [_old(i, f"part{i} words here", 3) for i in range(3)]
    new = [_new(0, "part0 words here\npart1 words here\npart2 words here",
                ["part0 words here", "part1 words here", "part2 words here"])]
    p = plan(old, new)
    assert sorted(p.deletes) == ["gdrive-F-1", "gdrive-F-2"]
    assert set(p.remap.values()) == {"gdrive-F-0"}
    assert all(p.reasons[k] == "exact" for k in p.remap)


def test_more_new_chunks_splits_and_remaps_by_start():
    old = [_old(0, "a1 a2 a3 a4 b1 b2 b3 b4", 1)]
    new = [_new(0, "a1 a2 a3 a4"), _new(1, "b1 b2 b3 b4")]
    p = plan(old, new)
    assert p.remap == {"gdrive-F-0": "gdrive-F-0"}
    assert p.deletes == [] and all(r.covered and r.enriched for r in p.rows)


def test_plan_repeated_text_maps_monotonically():
    boiler = "Confidential disclaimer text"
    old = [_old(0, f"{boiler} first", 2), _old(1, f"{boiler} second", 2)]
    new = [_new(0, f"{boiler} first"), _new(1, f"{boiler} second")]
    p = plan(old, new)
    assert p.remap == {"gdrive-F-0": "gdrive-F-0", "gdrive-F-1": "gdrive-F-1"}


def test_removed_repeated_header_maps_nearest():
    old = [_old(0, "Header one alpha", 3), _old(1, "Header", 3), _old(2, "gamma delta", 3)]
    new = [_new(0, "Header one alpha"), _new(1, "gamma delta")]
    p = plan(old, new)
    assert p.remap["gdrive-F-2"] == "gdrive-F-1"
    assert p.remap["gdrive-F-1"] in {"gdrive-F-0", "gdrive-F-1"}
    assert p.deletes == ["gdrive-F-2"]
    # every old id maps to a new chunk that exists
    assert set(p.remap.values()) <= {"gdrive-F-0", "gdrive-F-1"}


def test_plan_gap_in_old_chunks_is_safe():
    old = [_old(0, "alpha beta", 4), _old(2, "epsilon zeta", 4)]      # 1 and 3 missing
    new = [_new(0, "alpha beta"), _new(1, "gamma delta"), _new(2, "epsilon zeta")]
    p = plan(old, new)
    assert [r.covered for r in p.rows] == [True, False, True]
    assert p.remap["gdrive-F-2"] == "gdrive-F-2"


def test_old_order_is_chunk_index_not_list_order():
    old = [_old(1, "gamma delta", 2), _old(0, "alpha beta", 2)]
    p = plan(old, [_new(0, "alpha beta"), _new(1, "gamma delta")])
    assert p.content_equal
    assert p.remap == {"gdrive-F-0": "gdrive-F-0", "gdrive-F-1": "gdrive-F-1"}


def test_table_chunk_covered_by_cell_values():
    old = [_old(0, "Item Cost Chairs 120", 1)]
    new = [_new(0, "Table (in Budget)\nItem: Chairs; Cost: 120", ["Item", "Cost", "Chairs", "120"])]
    assert plan(old, new).rows[0].covered


def test_table_cell_matches_row_sentence_punctuation():
    old = [_old(0, "Item: Chairs; Cost: 120", 1)]
    new = [_new(0, "Item | Cost\nChairs | 120", ["Item", "Cost", "Chairs", "120"])]
    assert plan(old, new).rows[0].covered


def test_largest_overlap_old_supplies_state():
    old = [_old(0, "a b c d e f g h", 2, state="cold"), _old(1, "i j", 2, state=None)]
    new = [_new(0, "a b c d e f g h i j")]
    p = plan(old, new)
    assert p.rows[0].enrich_state == "cold"


def test_empty_string_state_reads_as_hot():
    p = plan([_old(0, "alpha beta", 1, state="")], [_new(0, "alpha beta")])
    assert p.rows[0].enrich_state is None


def test_lineage_gone_maps_to_first_new_chunk():
    old = [{**_old(0, "body text", 1), "doc_id": "gmail-M-body-0",
            "metadata": {"source_type": "gmail", "chunk_index": 0}},
           {**_old(0, "attachment words", 1), "doc_id": "gmail-M-att-0-0",
            "metadata": {"source_type": "gmail", "chunk_index": 0}}]
    new = [Chunk("gmail-M-body-0", "body text", "h", {"source_type": "gmail"}, ["body text"])]
    p = plan(old, new)
    assert p.remap["gmail-M-att-0-0"] == "gmail-M-body-0"
    assert p.reasons["gmail-M-att-0-0"] == "lineage_gone"
    assert p.unequal == ["gmail-M-att-0"] and not p.content_equal
    assert p.deletes == ["gmail-M-att-0-0"]


def test_no_new_chunks_refuses():
    import pytest
    with pytest.raises(ValueError):
        plan([_old(0, "alpha", 1)], [])


def test_norm():
    assert norm("  a\n\tb  c ") == "a b c"


def _deck_old(n):
    return [_old(i, f"Slide {i} title words s{i}a s{i}b", n) for i in range(n)]


def test_partly_located_chunks_stay_remap_targets():
    """Every new chunk = its old slide + new speaker notes: none is covered, but
    each still sits where its slide's text is, so the remap stays i -> i."""
    old = _deck_old(4)
    new = [_new(i, f"{o['text']}\nNotes: fresh remark n{i}",
                [o["text"], f"Notes: fresh remark n{i}"]) for i, o in enumerate(old)]
    p = plan(old, new)
    assert not any(r.covered for r in p.rows)
    assert all(r.enriched == 0 for r in p.rows)
    assert p.remap == {f"gdrive-F-{i}": f"gdrive-F-{i}" for i in range(4)}
    assert set(p.reasons.values()) == {"exact"}


def test_notes_on_odd_slides_only_keep_positions():
    old = _deck_old(4)
    new = []
    for i, o in enumerate(old):
        if i % 2:
            new.append(_new(i, f"{o['text']}\nNotes: n{i}", [o["text"], f"Notes: n{i}"]))
        else:
            new.append(_new(i, o["text"]))
    p = plan(old, new)
    assert [r.covered for r in p.rows] == [True, False, True, False]
    assert p.remap["gdrive-F-1"] == "gdrive-F-1"
    assert p.remap["gdrive-F-3"] == "gdrive-F-3"
    assert p.remap == {f"gdrive-F-{i}": f"gdrive-F-{i}" for i in range(4)}
    assert set(p.reasons.values()) == {"exact"}


# -- final review C1: remap on the UNION of a new chunk's span intervals -------

def _moved_table_docx():
    """A DOCX with a mid-document table, chunked the way the OLD extractor did
    (paragraphs '\n'-joined, every table appended at the END) and the way the
    new block extractor does (table in place)."""
    import io
    import random

    from docx import Document

    from mcpbrain.sync.blocks import render
    from mcpbrain.sync.extract_office import extract_blocks_from_docx
    from tests.oracles.chunking_v0 import chunk_text_v0
    rnd = random.Random(1)
    words = ("alpha beta gamma delta budget review staff campus report minutes "
             "action item follow the plan approved").split()

    def sent(n):
        return " ".join(rnd.choice(words) for _ in range(n)).capitalize() + "."
    doc = Document()
    doc.add_heading("Annual Review", 1)
    for s in range(6):
        doc.add_heading(f"Section {s}", 2)
        for _ in range(5):
            doc.add_paragraph(" ".join(sent(rnd.randint(8, 20)) for _ in range(4)))
        if s == 2:
            t = doc.add_table(rows=4, cols=3)
            for r in range(4):
                for c in range(3):
                    t.cell(r, c).text = f"cell{r}{c} {rnd.choice(words)}"
    bio = io.BytesIO()
    doc.save(bio)
    data = bio.getvalue()
    d2 = Document(io.BytesIO(data))
    parts = [p.text for p in d2.paragraphs if p.text.strip()]
    for table in d2.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    old_texts = chunk_text_v0("\n".join(parts))
    new = [Chunk(f"gdrive-F-{i}", r.text, "h", {"source_type": "gdrive", "file_id": "F",
                                                   "chunk_index": i}, r.spans)
           for i, r in enumerate(render(extract_blocks_from_docx(data)))]
    return [_old(i, t, len(old_texts)) for i, t in enumerate(old_texts)], new


def test_moved_table_docx_maps_each_old_id_to_the_chunk_holding_most_of_its_text():
    import re
    old, new = _moved_table_docx()
    p = plan(old, new)
    new_norm = {c.doc_id: norm(c.text) for c in new}
    for o in old:
        # ground truth: which new chunk holds the most of this old chunk's
        # sentences (random sentences are unique in this fixture)
        sents = [s for s in re.split(r"(?<=\.)\s+", norm(o["text"])) if len(s) > 30]
        score = {d: sum(len(s) for s in sents if s in t) for d, t in new_norm.items()}
        best = max(score.values())
        assert best > 0
        winners = {d for d, v in score.items() if v == best}
        assert p.remap[o["doc_id"]] in winners, (o["doc_id"], p.remap[o["doc_id"]], score)
    # the reviewer's probe collapsed old 4..7 onto new 3; they must spread out
    assert len({p.remap[o["doc_id"]] for o in old[4:]}) > 1


def test_remap_prefers_largest_union_overlap_over_first_hull():
    """New chunk 0 holds text from old 0 AND a far-later old 2 (a moved
    table); its hull swallows old 1, which must still map to new 1 where its
    text now lives."""
    old = [_old(0, "aa ab ac ad", 3), _old(1, "ba bb bc bd", 3), _old(2, "ca cb cc cd", 3)]
    new = [_new(0, "aa ab ac ad\nca cb cc cd", ["aa ab ac ad", "ca cb cc cd"]),
           _new(1, "ba bb bc bd")]
    p = plan(old, new)
    assert p.remap == {"gdrive-F-0": "gdrive-F-0", "gdrive-F-1": "gdrive-F-1",
                       "gdrive-F-2": "gdrive-F-0"}
    assert set(p.reasons.values()) == {"exact"}
    # coverage uses the union too: new 0 overlaps old 0 and old 2, not old 1
    old[1]["enriched"] = 0
    p = plan(old, new)
    assert p.rows[0].covered and p.rows[0].enriched == 1
    assert p.rows[1].covered and p.rows[1].enriched == 0


def test_coincidental_later_span_does_not_advance_the_cursor():
    """Task 9 triage: a partial chunk whose only located span is a short word
    that coincidentally occurs far later must not push the forward cursor past
    repeated text -- the later identical chunks still map in order."""
    boiler = "Confidential boiler text"
    old = [_old(0, "Intro words here", 5), _old(1, f"{boiler} one", 5),
           _old(2, "middle words", 5), _old(3, f"{boiler} one", 5), _old(4, "closing", 5)]
    new = [_new(0, "Intro words here"),
           _new(1, "Fresh notes\ncl", ["Fresh notes", "closing"]),   # coincidental hit
           _new(2, f"{boiler} one"), _new(3, "middle words"),
           _new(4, f"{boiler} one"), _new(5, "closing")]
    p = plan(old, new)
    assert p.remap["gdrive-F-1"] == "gdrive-F-2"
    assert p.remap["gdrive-F-3"] == "gdrive-F-4"
    assert p.reasons["gdrive-F-1"] == p.reasons["gdrive-F-3"] == "exact"


def test_repeated_header_partial_maps_to_its_body():
    """Task 9 triage: old chunks each began with a page header the new
    extractor emits once; the old chunk's start sits on a header copy no new
    chunk claims, but most of its text is in one new chunk -- map it there,
    exactly, not to a neighbour by distance."""
    hdr = "Northgate Trust board pack"
    old = [_old(0, f"{hdr} alpha one two three", 3),
           _old(1, f"{hdr} beta four five six seven eight", 3),
           _old(2, f"{hdr} gamma nine ten eleven twelve", 3)]
    new = [_new(0, f"{hdr}\nalpha one two three", [hdr, "alpha one two three"]),
           _new(1, "beta four five six seven eight"),
           _new(2, "gamma nine ten eleven twelve")]
    p = plan(old, new)
    assert p.remap == {"gdrive-F-0": "gdrive-F-0", "gdrive-F-1": "gdrive-F-1",
                       "gdrive-F-2": "gdrive-F-2"}
    assert set(p.reasons.values()) == {"exact"}


# -- final review: ONE pending-unit scanner, the handler's superset of keys ----

def test_pending_unit_refs_covers_every_id_key_and_rescans_on_change(tmp_path):
    import json
    import os
    import time

    from mcpbrain import reflow
    q = tmp_path / "enrich_queue"
    (q / "units").mkdir(parents=True)
    (q / "units" / "u1.json").write_text(json.dumps({"threads": [
        {"thread_id": "T", "doc_id": "D", "file_id": "FI", "event_id": "EV",
         "session_id": "SE", "doc_ids": ["X0", {"doc_id": "X1"}],
         "messages": [{"message_id": "M", "part_doc_ids": ["P0"],
                       "chunk_doc_ids": ["C0"]}]}]}))
    assert reflow.pending_unit_refs(tmp_path) >= {
        "T", "D", "FI", "EV", "SE", "X0", "X1", "M", "P0", "C0"}
    (q / "units" / "u2.json").write_text(json.dumps({"thread_id": "NEW"}))
    st = (q / "units").stat()
    os.utime(q / "units", ns=(st.st_atime_ns, st.st_mtime_ns + 10_000_000))
    time.sleep(0)
    assert "NEW" in reflow.pending_unit_refs(tmp_path)


def test_handler_uses_the_shared_scanner():
    from mcpbrain.sync import reflow_handler
    assert not hasattr(reflow_handler, "_collect_refs")
    assert not hasattr(reflow_handler.ReflowContext, "_pending_unit_refs")


def test_duplicate_old_text_maps_to_the_new_chunk_holding_it():
    """Dry run #2 D1/D2: an old chunk whose text duplicates another old
    chunk's (a legacy positional tail) gets no placed new chunk on its own
    region; it maps 'exact' onto the new chunk holding that text, never
    'nearest' by position."""
    old = [_old(0, "alpha beta gamma", 2), _old(1, "delta epsilon zeta", 2),
           _old(2, "alpha beta", 2)]
    p = plan(old, [_new(0, "alpha beta gamma"), _new(1, "delta epsilon zeta")])
    assert p.remap["gdrive-F-2"] == "gdrive-F-0"
    assert p.reasons["gdrive-F-2"] == "exact"
    assert p.deletes == ["gdrive-F-2"]
