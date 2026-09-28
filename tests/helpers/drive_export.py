"""Make a fake Drive `files().export` behave like the real one: convert the
stored document body to the REQUESTED format.

Since 2026-09-24 Google Docs/Slides export as DOCX/PPTX (drive._EXPORT_BLOCKS)
rather than text/plain. Fakes that store a plain-text body per fileId pass it
through `as_export(body, mimeType)` so the same fixtures keep working and keep
their assertions about the resulting text.
"""
import io

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def as_export(payload, mime_type: str):
    if not isinstance(payload, (bytes, str)) or not payload:
        return payload
    raw = payload if isinstance(payload, bytes) else payload.encode("utf-8")
    if raw.startswith(b"PK"):            # already an OOXML zip
        return raw
    text = raw.decode("utf-8", "replace")
    paras = [p for p in text.split("\n\n") if p.strip()]
    if mime_type == DOCX:
        from docx import Document
        d = Document()
        for p in paras:
            d.add_paragraph(p)
        b = io.BytesIO(); d.save(b); return b.getvalue()
    if mime_type == PPTX:
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(4)
                             ).text_frame.text = text
        b = io.BytesIO(); prs.save(b); return b.getvalue()
    return payload
