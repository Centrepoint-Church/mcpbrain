"""Task 7 (2026-09-24 extraction-fidelity): block extractors wired into Drive,
Gmail attachments and every normaliser; version stamps; heading trail."""
import io

from docx import Document

from mcpbrain.chunking import SPLIT_VERSION
from mcpbrain.embed import contextual_prefix
from mcpbrain.sync import drive
from mcpbrain.sync.blocks import Heading, Paragraph
from mcpbrain.sync.normalise import normalise_gmail

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
GDOC = "application/vnd.google-apps.document"
GSLIDES = "application/vnd.google-apps.presentation"


def _docx_bytes():
    d = Document()
    d.add_heading("Budget", 1)
    d.add_paragraph("Line one\nLine two")
    b = io.BytesIO(); d.save(b); return b.getvalue()


def _pptx_bytes():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)
                         ).text_frame.text = "Quarterly roster for Northgate Trust"
    b = io.BytesIO(); prs.save(b); return b.getvalue()


class _Req:
    def __init__(self, payload=None, exc=None):
        self.payload, self.exc = payload, exc
    def execute(self, num_retries=0):
        if self.exc:
            raise self.exc
        return self.payload


class _Files:
    def __init__(self, media=None, exports=None):
        self.media, self.exports, self.export_calls = media, exports or {}, []
    def get_media(self, fileId, supportsAllDrives=True):
        return _Req(self.media)
    def export(self, fileId, mimeType):
        self.export_calls.append(mimeType)
        val = self.exports.get(mimeType)
        return _Req(exc=val) if isinstance(val, Exception) else _Req(val)


class _Svc:
    def __init__(self, files):
        self._f = files
    def files(self):
        return self._f


def _size_error():
    from googleapiclient.errors import HttpError

    class _Resp(dict):
        status = 403
        reason = "exportSizeLimitExceeded"
    return HttpError(_Resp(), b'{"error": {"errors": [{"reason": "exportSizeLimitExceeded"}]}}')


def test_docx_drive_file_uses_blocks_and_stamps_versions():
    svc = _Svc(_Files(media=_docx_bytes()))
    meta = {"id": "f1", "name": "b.docx", "mimeType": DOCX, "modifiedTime": "2026-01-01T00:00:00Z"}
    content = drive.fetch_content(svc, meta)
    assert content.blocks and isinstance(content.blocks[0], Heading)
    chunks = drive.normalise_drive(meta, content.text, blocks=content.blocks)
    assert chunks[0].text == "Budget\n\nLine one\nLine two"
    md = chunks[0].metadata
    assert md["extraction_version"] == 1 and md["split_version"] == SPLIT_VERSION
    assert md["heading_trail"] == "Budget"
    assert chunks[0].spans == ["Budget", "Line one\nLine two"]


def test_google_doc_exports_docx_then_falls_back_on_size_error():
    files = _Files(exports={DOCX: _size_error(), "text/markdown": b"# Title\n\nBody text"})
    meta = {"id": "g1", "name": "Doc", "mimeType": GDOC, "modifiedTime": "2026-01-01T00:00:00Z"}
    content = drive.fetch_content(_Svc(files), meta)
    assert files.export_calls == [DOCX, "text/markdown"]
    assert content.blocks == [Heading(1, "Title"), Paragraph("Body text")]


def test_google_doc_falls_back_to_plain_text_after_markdown_size_error():
    files = _Files(exports={DOCX: _size_error(), "text/markdown": _size_error(),
                            "text/plain": b"Plain body"})
    meta = {"id": "g1", "name": "Doc", "mimeType": GDOC, "modifiedTime": "2026-01-01T00:00:00Z"}
    content = drive.fetch_content(_Svc(files), meta)
    assert files.export_calls == [DOCX, "text/markdown", "text/plain"]
    assert content.text == "Plain body"


def test_google_doc_non_size_export_error_raises_not_falls_back():
    """Only the size cap falls back; anything else (a 500, an auth error) must
    raise so work_queue backs the item off rather than silently degrading."""
    import pytest
    boom = RuntimeError("backend error")
    files = _Files(exports={DOCX: boom, "text/markdown": b"# never"})
    meta = {"id": "g1", "name": "Doc", "mimeType": GDOC}
    with pytest.raises(RuntimeError):
        drive.fetch_content(_Svc(files), meta)
    assert files.export_calls == [DOCX]


def test_google_doc_default_path_is_docx():
    files = _Files(exports={DOCX: _docx_bytes()})
    meta = {"id": "g1", "name": "Doc", "mimeType": GDOC, "modifiedTime": "2026-01-01T00:00:00Z"}
    content = drive.fetch_content(_Svc(files), meta)
    assert files.export_calls == [DOCX]
    chunks = drive.normalise_drive(meta, content.text, blocks=content.blocks)
    assert chunks[0].metadata["extraction_version"] == 1
    assert chunks[0].metadata["heading_trail"] == "Budget"


def test_google_slides_export_pptx_then_plain_text():
    files = _Files(exports={PPTX: _pptx_bytes()})
    meta = {"id": "s1", "name": "Deck", "mimeType": GSLIDES}
    content = drive.fetch_content(_Svc(files), meta)
    assert files.export_calls == [PPTX]
    assert isinstance(content.blocks[0], Heading) and content.blocks[0].text == "Slide 1"
    assert "Quarterly roster for Northgate Trust" in content.text

    files = _Files(exports={PPTX: _size_error(), "text/plain": b"Slide text"})
    content = drive.fetch_content(_Svc(files), meta)
    assert files.export_calls == [PPTX, "text/plain"]
    assert content.text == "Slide text"


def test_rtf_drive_file_goes_through_decoder_not_verbatim():
    rtf = rb"{\rtf1\ansi{\fonttbl{\f0 Arial;}}Hello \'e9t\'e9\par Second para}"
    meta = {"id": "r1", "name": "n.rtf", "mimeType": "application/rtf",
            "modifiedTime": "2026-01-01T00:00:00Z"}
    content = drive.fetch_content(_Svc(_Files(media=rtf)), meta)
    assert "\\rtf1" not in content.text and "Arial" not in content.text
    assert "Hello été" in content.text and "Second para" in content.text
    chunks = drive.normalise_drive(meta, content.text, blocks=content.blocks)
    assert chunks[0].metadata["extraction_version"] == 1


def test_extract_blocks_from_rtf_happy_path():
    from mcpbrain.sync.extractors import extract_blocks_from_rtf
    rtf = (rb"{\rtf1\ansi\ansicpg1252{\fonttbl{\f0 Arial;}}{\colortbl;\red0\green0\blue0;}"
           rb"{\*\generator Word;}Caf\'e9 plans\par\par Na\u239?ve second paragraph\line "
           rb"with a line break\par}")
    assert extract_blocks_from_rtf(rtf) == [
        Paragraph("Café plans"),
        Paragraph("Naïve second paragraph\nwith a line break"),
    ]


def test_extract_blocks_from_rtf_bad_input_is_empty_not_raising(monkeypatch):
    from mcpbrain.sync import rtf as rtf_mod
    monkeypatch.setattr(rtf_mod, "rtf_to_text", lambda _d: (_ for _ in ()).throw(ValueError("x")))
    assert rtf_mod.extract_blocks_from_rtf(b"{\\rtf1 x}") == []


def test_prose_drive_file_still_stamps_split_version_but_no_extraction_version():
    meta = {"id": "t1", "name": "n.txt", "mimeType": "text/plain"}
    chunks = drive.normalise_drive(meta, "Plain note body text")
    md = chunks[0].metadata
    assert md["split_version"] == SPLIT_VERSION
    assert "extraction_version" not in md and "heading_trail" not in md
    assert chunks[0].spans == ["Plain note body text"]


def test_table_drive_chunks_stamp_split_version_and_spans():
    from mcpbrain.sync.tabular import Table
    meta = {"id": "x1", "name": "b.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    t = Table(sheet="S", header=["Item", "Cost"], rows=[["Chairs", "120"]],
              rows_total=1, truncated=False)
    chunks = drive.normalise_drive(meta, "", tables=[t])
    assert chunks and chunks[0].metadata["split_version"] == SPLIT_VERSION
    assert chunks[0].spans == [chunks[0].text]


def test_partial_block_extraction_is_marked_partial(monkeypatch):
    from mcpbrain.sync.blocks import PartialBlocks
    monkeypatch.setitem(drive._DOWNLOAD_BLOCKS, "application/pdf",
                        lambda _b: PartialBlocks([Paragraph("First page only")]))
    meta = {"id": "p1", "name": "r.pdf", "mimeType": "application/pdf"}
    content = drive.fetch_content(_Svc(_Files(media=b"%PDF")), meta)
    assert content.partial is True and content.text == "First page only"


def test_empty_block_extraction_is_empty_content(monkeypatch):
    monkeypatch.setitem(drive._DOWNLOAD_BLOCKS, "application/pdf", lambda _b: [])
    meta = {"id": "p1", "name": "r.pdf", "mimeType": "application/pdf"}
    report: dict = {}
    content = drive.fetch_content(_Svc(_Files(media=b"%PDF")), meta, report=report)
    assert content is not None and not content.text and not content.blocks
    assert report == {("extraction_empty", "application/pdf"): 1}


def test_handle_drive_item_threads_blocks_through(tmp_path):
    from mcpbrain.store import Store
    s = Store(tmp_path / "a.sqlite3", dim=4); s.init()

    class _F(_Files):
        def get(self, **kw):
            return _Req({"id": kw["fileId"], "name": "b.docx", "mimeType": DOCX,
                         "modifiedTime": "2026-01-01T00:00:00Z", "parents": []})
    drive.handle_drive_item(_Svc(_F(media=_docx_bytes())), s,
                            {"ref_id": "f1", "event": "upsert"})
    rows = s.owner_chunks(["gdrive-f1-"])
    assert rows and rows[0]["metadata"]["heading_trail"] == "Budget"
    assert rows[0]["metadata"]["extraction_version"] == 1


def test_cache_first_extract_passes_mime_to_both_try_imports(monkeypatch):
    from mcpbrain import ingest_cache
    seen = []
    monkeypatch.setattr(ingest_cache, "try_import",
                        lambda *a, **k: seen.append(k.get("mime")) or False)
    monkeypatch.setattr(drive, "upsert_file_chunks", lambda *a, **k: None)
    monkeypatch.setattr(drive, "folder_path", lambda *a, **k: "")
    meta = {"id": "f1", "name": "b.docx", "mimeType": DOCX, "modifiedTime": "2026-01-01T00:00:00Z"}
    ok, miss = drive._cache_first_extract_one(
        _Svc(_Files(media=_docx_bytes())), object(), object(), "D1", meta, {})
    assert ok and seen == [DOCX, DOCX]


def test_gmail_body_chunks_carry_split_version():
    import base64
    body = base64.urlsafe_b64encode(b"Hello there\n\nSecond para").decode()
    raw = {"id": "m1", "threadId": "t1", "labelIds": [],
           "payload": {"mimeType": "text/plain", "headers": [{"name": "Subject", "value": "Hi"}],
                       "body": {"data": body}}}
    chunks = normalise_gmail(raw)
    assert chunks and all(c.metadata["split_version"] == SPLIT_VERSION for c in chunks)


def test_calendar_and_anarlog_chunks_carry_split_version():
    from mcpbrain.sync.anarlog import normalise_session
    from mcpbrain.sync.calendar import normalise_calendar
    ev = {"id": "e1", "summary": "Staff meeting", "description": "Agenda items",
          "start": {"dateTime": "2026-01-01T09:00:00Z"}, "end": {"dateTime": "2026-01-01T10:00:00Z"}}
    cal = normalise_calendar(ev)
    assert cal and all(c.metadata["split_version"] == SPLIT_VERSION for c in cal)
    sess = {"id": "S1", "title": "Staff meeting", "created_at": "2026-01-01T09:00:00Z",
            "documents": {}, "transcript": "Dana Okafor: we agreed the roster."}
    an = normalise_session(sess)
    assert an and all(c.metadata["split_version"] == SPLIT_VERSION for c in an)


def test_pdf_attachment_uses_blocks_and_stamps_versions(monkeypatch):
    from mcpbrain.sync import attachments
    monkeypatch.setitem(attachments._BLOCK_EXTRACTORS, "application/pdf",
                        lambda _b: [Heading(1, "Invoice"), Paragraph("Chairs 120\nTables 80")])
    raw = {"id": "M", "threadId": "T", "payload": {"headers": []}}
    part = {"filename": "i.pdf", "mime": "application/pdf", "index": 0}
    chunks = attachments.normalise_attachment(raw, part, b"%PDF")
    md = chunks[0].metadata
    assert chunks[0].doc_id == "gmail-M-att-0-0"
    assert md["extraction_version"] == 1 and md["split_version"] == SPLIT_VERSION
    assert md["heading_trail"] == "Invoice"
    assert chunks[0].spans == ["Invoice", "Chairs 120\nTables 80"]


def test_docx_attachment_real_bytes_and_non_block_attachment_stamp():
    from mcpbrain.sync import attachments
    raw = {"id": "M", "threadId": "T", "payload": {"headers": []}}
    chunks = attachments.normalise_attachment(
        raw, {"filename": "b.docx", "mime": DOCX, "index": 1}, _docx_bytes())
    assert chunks[0].text == "Budget\n\nLine one\nLine two"
    assert chunks[0].metadata["extraction_version"] == 1
    txt = attachments.normalise_attachment(
        raw, {"filename": "n.txt", "mime": "text/plain", "index": 2}, b"Just a note")
    assert txt[0].metadata["split_version"] == SPLIT_VERSION
    assert "extraction_version" not in txt[0].metadata


def test_rtf_attachment_is_supported():
    from mcpbrain.sync import attachments
    assert attachments._supported("application/rtf")
    raw = {"id": "M", "threadId": "T", "payload": {"headers": []}}
    chunks = attachments.normalise_attachment(
        raw, {"filename": "n.rtf", "mime": "application/rtf", "index": 0},
        rb"{\rtf1\ansi Plan for the week\par}")
    assert chunks[0].text == "Plan for the week"
    assert chunks[0].metadata["extraction_version"] == 1


def test_contextual_prefix_includes_heading_trail():
    p = contextual_prefix({"source_type": "gdrive", "file_name": "b.docx",
                           "heading_trail": "Budget › Capital works"})
    assert "section Budget › Capital works" in p


def test_contextual_prefix_without_trail_unchanged():
    p = contextual_prefix({"source_type": "gdrive", "file_name": "b.docx"})
    assert "section" not in p
