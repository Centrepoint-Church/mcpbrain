import io

from pptx import Presentation
from pptx.util import Inches

from mcpbrain.sync.blocks import Heading, Paragraph, TableBlock
from mcpbrain.sync.extract_office import extract_blocks_from_pptx


def _pptx() -> bytes:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])          # title only
    s.shapes.title.text = "Q3 Ministry Review"
    group = s.shapes.add_group_shape()
    tb = group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    tb.text_frame.text = "Inside a group"
    rows = s.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    rows.cell(0, 0).text, rows.cell(0, 1).text = "Campus", "Attendance"
    rows.cell(1, 0).text, rows.cell(1, 1).text = "North", "140"
    s.notes_slide.notes_text_frame.text = "Mention the Northgate Trust grant."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_slide_heading_group_table_and_notes():
    blocks = extract_blocks_from_pptx(_pptx())
    assert blocks[0] == Heading(2, "Slide 1: Q3 Ministry Review")
    assert Paragraph("Inside a group") in blocks
    assert TableBlock([["Campus", "Attendance"], ["North", "140"]]) in blocks
    assert blocks[-1] == Paragraph("Notes: Mention the Northgate Trust grant.")


def test_garbage_input_returns_empty_list():
    assert extract_blocks_from_pptx(b"not a pptx") == []
