import io

from docx import Document

from mcpbrain.sync.blocks import Heading, Paragraph, TableBlock
from mcpbrain.sync.extract_office import extract_blocks_from_docx


def _docx() -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Northgate Trust — Confidential"
    doc.sections[0].footer.paragraphs[0].text = "Page footer"
    doc.add_heading("Quarterly review", level=1)
    doc.add_paragraph("Intro paragraph before the table.")
    t = doc.add_table(rows=3, cols=3)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(0, 2).text = "Item", "Owner", "Due"
    t.cell(1, 0).text, t.cell(1, 1).text = "Roster", "Dana Okafor"
    t.cell(2, 0).text = "Merged"
    t.cell(2, 1).merge(t.cell(2, 2)).text = "Spans two"
    doc.add_heading("Next steps", level=2)
    doc.add_paragraph("After the table.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_body_order_headings_and_table_in_place():
    blocks = extract_blocks_from_docx(_docx())
    kinds = [type(b).__name__ for b in blocks]
    assert blocks[0] == Paragraph("Northgate Trust — Confidential")
    assert Heading(1, "Quarterly review") in blocks
    ti = next(i for i, b in enumerate(blocks) if isinstance(b, TableBlock))
    intro = blocks.index(Paragraph("Intro paragraph before the table."))
    after = blocks.index(Paragraph("After the table."))
    assert intro < ti < after
    assert Heading(2, "Next steps") in blocks
    assert blocks[-1] == Paragraph("Page footer")
    assert kinds.count("TableBlock") == 1


def test_empty_cells_kept_and_merged_cells_once():
    t = next(b for b in extract_blocks_from_docx(_docx()) if isinstance(b, TableBlock))
    assert t.rows[1] == ["Roster", "Dana Okafor", ""]
    assert t.rows[2] == ["Merged", "Spans two"]


def test_garbage_input_returns_empty_list():
    assert extract_blocks_from_docx(b"nope") == []


def _docx_with_text_box_and_vmerge() -> bytes:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    doc = Document()
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).merge(t.cell(1, 0)).text = "Campus North"
    t.cell(0, 1).text, t.cell(1, 1).text = "Sunday", "Wednesday"
    t.cell(2, 0).text, t.cell(2, 1).text = "South", "Friday"
    p = doc.add_paragraph("Anchor paragraph.")
    p._p.append(parse_xml(
        '<w:r %s><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/'
        'drawingml/2006/wordprocessingDrawing"><a:graphic xmlns:a="http://schemas.'
        'openxmlformats.org/drawingml/2006/main"><a:graphicData uri="x"><wps:wsp '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:txbx><w:txbxContent><w:p><w:r><w:t>Box line one</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Box line two</w:t></w:r></w:p></w:txbxContent></wps:txbx>'
        '</wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
        % nsdecls("w")))
    doc.add_paragraph("After the box.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_text_box_follows_its_anchor_paragraph():
    blocks = extract_blocks_from_docx(_docx_with_text_box_and_vmerge())
    i = blocks.index(Paragraph("Anchor paragraph."))
    assert blocks[i + 1] == Paragraph("Box line one\nBox line two")
    assert blocks[i + 2] == Paragraph("After the box.")


def test_vertically_merged_cell_repeats_per_row():
    """Pins extraction_version 1: python-docx resolves each row's vertical-merge
    continuation to the origin cell, so its text repeats on every row it spans
    (horizontal merges are de-duplicated within a row). Changing this changes chunk text."""
    t = next(b for b in extract_blocks_from_docx(_docx_with_text_box_and_vmerge())
             if isinstance(b, TableBlock))
    assert t.rows == [["Campus North", "Sunday"], ["Campus North", "Wednesday"],
                      ["South", "Friday"]]
